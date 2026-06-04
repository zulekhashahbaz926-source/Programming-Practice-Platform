from pptx import Presentation
from pptx.util import Inches, Pt
from pathlib import Path
import os

OUT = Path('docs')
OUT.mkdir(exist_ok=True)
PPTX_PATH = OUT / 'project_presentation.pptx'

prs = Presentation()
# Title slide
slide_layout = prs.slide_layouts[0]
slide = prs.slides.add_slide(slide_layout)
slide.shapes.title.text = "Zyntriva CodeForge — Project Presentation"
slide.placeholders[1].text = "High-level overview, microservices, and diagrams"

# Overview slide
slide_layout = prs.slide_layouts[1]
slide = prs.slides.add_slide(slide_layout)
slide.shapes.title.text = "Project Overview"
body = slide.placeholders[1].text_frame
body.clear()
lines = [
    "Python desktop app with CustomTkinter",
    "Modules: problems, quizzes, testing, peer review",
    "SQLite persistence",
    "Microservices: Weather, Notification, Analytics, Preferences",
]
for l in lines:
    p = body.add_paragraph()
    p.text = l
    p.level = 0
    p.font.size = Pt(18)

# Microservices slide
slide = prs.slides.add_slide(slide_layout)
slide.shapes.title.text = "Microservices"
body = slide.placeholders[1].text_frame
body.clear()
services = [
    "Weather API — microservices/weather_service/app.py",
    "Notification — microservices/notification_service/app.py",
    "Analytics — microservices/analytics_service/app.py",
    "Preferences — microservices/preferences_service/app.py",
]
for s in services:
    p = body.add_paragraph()
    p.text = s
    p.level = 0
    p.font.size = Pt(16)

# Add any images in docs/uml
uml_dir = Path('docs/uml')
if uml_dir.exists():
    for img in sorted(uml_dir.glob('*.png')):
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = img.stem.replace('_', ' ').title()
        slide.shapes.add_picture(str(img), Inches(1), Inches(1.5), width=Inches(8))

prs.save(PPTX_PATH)
print('Saved presentation to', PPTX_PATH)
