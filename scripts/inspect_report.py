from pathlib import Path
from docx import Document
from pptx import Presentation

base = Path(__file__).resolve().parent.parent
report_doc = base / 'docs' / 'project_report.docx'
presentation_file = base / 'docs' / 'project_presentation.pptx'

print('DOCX exists:', report_doc.exists())
if report_doc.exists():
    doc = Document(report_doc)
    print('DOCX paragraphs:', len(doc.paragraphs))
    for i, p in enumerate(doc.paragraphs[:40]):
        print(i, repr(p.text))

print('PPTX exists:', presentation_file.exists())
if presentation_file.exists():
    prs = Presentation(presentation_file)
    print('PPTX slides:', len(prs.slides))
    for i, slide in enumerate(prs.slides):
        print('Slide', i, 'title:', slide.shapes.title.text if slide.shapes.title else None)
        for shape in slide.shapes:
            if shape.has_text_frame:
                print('  text:', shape.text)
