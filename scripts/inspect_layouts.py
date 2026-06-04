from pptx import Presentation
from pathlib import Path
prs = Presentation()
for i, layout in enumerate(prs.slide_layouts):
    print('Layout', i)
    for shape in layout.shapes:
        print('  shape', type(shape), 'name', shape.name, 'has_text', shape.has_text_frame, 'placeholder type', getattr(shape, 'placeholder_format', None))
