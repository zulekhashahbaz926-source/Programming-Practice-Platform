from pathlib import Path
from docx import Document
from pptx import Presentation
from pptx.util import Inches, Pt

BASE_DIR = Path(__file__).resolve().parent.parent
DOCS_DIR = BASE_DIR / "docs"
OUTPUT_DOCX = DOCS_DIR / "project_report.docx"
OUTPUT_PPTX = DOCS_DIR / "project_presentation.pptx"
REPORT_MD = DOCS_DIR / "semester_project_report.md"
UML_IMAGES = {
    "Use Case": DOCS_DIR / "uml_use_case.png",
    "Class": DOCS_DIR / "uml_class_diagram.png",
    "Activity": DOCS_DIR / "uml_activity_diagram.png",
    "Sequence": DOCS_DIR / "uml_sequence_diagram.png",
}


def load_markdown_sections(path: Path):
    sections = []
    current_title = None
    current_lines = []
    with path.open("r", encoding="utf-8") as f:
        for line in f:
            stripped = line.strip()
            if stripped.startswith("## "):
                if current_title:
                    sections.append((current_title, current_lines))
                    current_lines = []
                current_title = stripped[3:]
            elif current_title:
                current_lines.append(line.rstrip())
    if current_title:
        sections.append((current_title, current_lines))
    return sections


def create_docx(sections):
    doc = Document()
    doc.add_heading("Zyntriva Semester Project Report", level=1)
    doc.add_paragraph(
        "This document summarizes the Zyntriva CodeForge desktop application and academic project deliverables."
    )
    for title, lines in sections:
        doc.add_heading(title, level=2)
        for line in lines:
            if line.startswith("-"):
                paragraph = doc.add_paragraph(line[1:].strip(), style="List Bullet")
            elif line:
                doc.add_paragraph(line)
    doc.add_page_break()
    doc.add_heading("High-Level UML Diagrams", level=2)
    for label, image_path in UML_IMAGES.items():
        if image_path.exists():
            doc.add_paragraph(label)
            doc.add_picture(str(image_path), width=Inches(6.5))
            doc.add_paragraph()
    doc.save(OUTPUT_DOCX)
    print(f"Saved report document to {OUTPUT_DOCX}")


def add_slide(prs, title, content_lines, image_path=None):
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = title
    text_frame = None
    for shape in slide.shapes:
        if shape.has_text_frame and shape.placeholder_format.type == 1:
            text_frame = shape.text_frame
            break
    if text_frame is None:
        text_frame = slide.shapes.add_textbox(Inches(0.5), Inches(1.8), Inches(8.0), Inches(4.5)).text_frame
    text_frame.clear()
    for line in content_lines:
        if not line:
            continue
        p = text_frame.add_paragraph()
        p.text = line
        p.font.size = Pt(18)
        p.level = 0
    if image_path and image_path.exists():
        slide.shapes.add_picture(str(image_path), Inches(5.5), Inches(1.5), width=Inches(4))


def create_pptx(sections):
    prs = Presentation()
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    slide.shapes.title.text = "Zyntriva Semester Project Presentation"
    slide.placeholders[1].text = (
        "A concise overview of the desktop application, UML architecture, testing, and software engineering project requirements."
    )
    for title, lines in sections[:6]:
        bullets = [line[2:].strip() for line in lines if line.startswith("- ")]
        add_slide(prs, title, bullets[:6])
    for label, image_path in UML_IMAGES.items():
        if image_path.exists():
            slide = prs.slides.add_slide(prs.slide_layouts[5])
            slide.shapes.title.text = f"{label} UML Diagram"
            slide.shapes.add_picture(str(image_path), Inches(1), Inches(1.5), width=Inches(8))
    prs.save(OUTPUT_PPTX)
    print(f"Saved presentation to {OUTPUT_PPTX}")


if __name__ == "__main__":
    sections = load_markdown_sections(REPORT_MD)
    create_docx(sections)
    create_pptx(sections)
