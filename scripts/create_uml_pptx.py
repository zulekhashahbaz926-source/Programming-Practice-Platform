from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pathlib import Path
import os

# Create presentation
prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(7.5)

# Slide 1: Title
slide1 = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
left = Inches(0.5)
top = Inches(2)
title_box = slide1.shapes.add_textbox(left, top, Inches(9), Inches(2))
title_frame = title_box.text_frame
title_frame.text = "Programming Practice Platform"
title_frame.paragraphs[0].font.size = Pt(54)
title_frame.paragraphs[0].font.bold = True
title_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
title_frame.paragraphs[0].font.color.rgb = RGBColor(30, 64, 175)

subtitle_box = slide1.shapes.add_textbox(left, Inches(4), Inches(9), Inches(1.5))
subtitle_frame = subtitle_box.text_frame
subtitle_frame.text = "UML Diagrams and Architecture"
subtitle_frame.paragraphs[0].font.size = Pt(32)
subtitle_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

# Slide 2: Class Diagram
slide2 = prs.slides.add_slide(prs.slide_layouts[1])
title2 = slide2.shapes.title
title2.text = "Class Diagram"
class_diagram_path = Path("docs/uml/class_diagram.png")
if class_diagram_path.exists():
    slide2.shapes.add_picture(str(class_diagram_path), Inches(0.5), Inches(1.5), width=Inches(9))

# Slide 3: Use Case Diagram
slide3 = prs.slides.add_slide(prs.slide_layouts[1])
title3 = slide3.shapes.title
title3.text = "Use Case Diagram"
usecase_path = Path("docs/uml/use_case_diagram.png")
if usecase_path.exists():
    slide3.shapes.add_picture(str(usecase_path), Inches(0.5), Inches(1.5), width=Inches(9))

# Slide 4: Sequence Diagram
slide4 = prs.slides.add_slide(prs.slide_layouts[1])
title4 = slide4.shapes.title
title4.text = "Sequence Diagram - Code Submission"
seq_path = Path("docs/uml/sequence_diagram.png")
if seq_path.exists():
    slide4.shapes.add_picture(str(seq_path), Inches(0.5), Inches(1.5), width=Inches(9))

# Slide 5: Activity Diagram
slide5 = prs.slides.add_slide(prs.slide_layouts[1])
title5 = slide5.shapes.title
title5.text = "Activity Diagram - Submit & Evaluate"
activity_path = Path("docs/uml/activity_diagram.png")
if activity_path.exists():
    slide5.shapes.add_picture(str(activity_path), Inches(0.5), Inches(1.5), width=Inches(9))

# Slide 6: Architecture Overview
slide6 = prs.slides.add_slide(prs.slide_layouts[1])
title6 = slide6.shapes.title
title6.text = "System Components"
body6 = slide6.placeholders[1]
text_frame6 = body6.text_frame
text_frame6.clear()

components = [
    ("Frontend (Tkinter GUI)", "User interface with problem browser, code editor, and dashboard"),
    ("Backend Services", "Authentication, evaluation engine, database management"),
    ("Database (SQLite)", "Stores users, problems, attempts, quizzes, and progress"),
    ("Evaluator Service", "Runs test cases against submitted code"),
    ("Microservices", "Weather API, Notifications, Analytics, User Preferences"),
]

for comp, desc in components:
    p = text_frame6.add_paragraph()
    p.text = f"{comp}: {desc}"
    p.level = 0
    p.font.size = Pt(16)

# Slide 7: Key Features
slide7 = prs.slides.add_slide(prs.slide_layouts[1])
title7 = slide7.shapes.title
title7.text = "Key Features"
body7 = slide7.placeholders[1]
text_frame7 = body7.text_frame
text_frame7.clear()

features = [
    "Multi-user authentication and authorization",
    "Code submission and automated evaluation",
    "Quiz management with time limits",
    "Progress tracking and badge system",
    "Admin dashboard for problem management",
    "Microservice architecture for scalability",
    "SQLite persistence layer",
    "Responsive Tkinter GUI with dark theme",
]

for feature in features:
    p = text_frame7.add_paragraph()
    p.text = feature
    p.level = 0
    p.font.size = Pt(18)

# Save presentation
output_path = Path("docs/uml/Programming_Practice_Platform_UML.pptx")
prs.save(str(output_path))
print(f"Created presentation: {output_path}")
